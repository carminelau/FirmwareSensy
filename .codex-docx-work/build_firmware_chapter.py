from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SOURCE = Path(r"C:\Users\carmi\OneDrive\Lavoro\1 - Sense Square 2021 -\SSQ - WORK\PKS Sense Square\Template_PKS_Sense_Square.docx")
OUTPUT = Path(r"C:\Work\FirmwareSensy\PKS_Sensy_Sezione_5_Firmware.docx")
GREEN = RGBColor(36, 126, 50)


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_cell_text(cell, label, text, font_size=9.0):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(font_size)
    r2 = p.add_run(text)
    r2.font.size = Pt(font_size)


def set_status(cell, symbol):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(symbol)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = GREEN


def fill_checklist(table, rows):
    assert len(table.rows) == len(rows)
    for row, (status, label, text) in zip(table.rows, rows):
        set_status(row.cells[0], status)
        set_cell_text(row.cells[1], label, text)


def clone_rows(table, wanted):
    while len(table.rows) < wanted:
        table._tbl.append(deepcopy(table.rows[-1]._tr))


def set_plain_cell(cell, text, font_size=8.0, bold=False, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(font_size)


doc = Document(SOURCE)

# Section 5 tables: 39 Repository, 40 Build/deploy, 41 Operations,
# 42 artifact register, 43 gaps. Source structure is retained.
fill_checklist(
    doc.tables[39],
    [
        (
            "☑",
            "Repository ufficiale. ",
            "https://github.com/carminelau/FirmwareSensy.git; branch osservato: main; baseline documentale: HEAD 7310c42 del 16/07/2026. La fotografia usata include modifiche locali non ancora consolidate. Owner nominale e matrice permessi: da approvare; responsabilità funzionale Software & AI.",
        ),
        (
            "◐",
            "Branch strategy. ",
            "Lo sviluppo corrente converge su main. Dashboard OTA e script di release possono compilare una branch selezionata tramite worktree temporaneo. Nel repository non risultano evidenze controllate di branch protection, pull request obbligatoria, reviewer o convenzione per branch: formalizzazione richiesta prima del rilascio commerciale.",
        ),
        (
            "◐",
            "Dipendenze e licenze. ",
            "platformio.ini blocca le dipendenze con versione o commit: RTClib, ArduinoJson, MQTT, ArduinoHttpClient, ESPAsyncWebServer, librerie Sensirion/DFRobot, GNSS, BH1750, Modbus e driver proprietari/fork. Riproducibilità buona; inventario licenze, notice e SBOM non presenti e quindi ancora da produrre.",
        ),
        (
            "☑",
            "Toolchain. ",
            "PlatformIO Core 6.1.19; piattaforma espressif32 6.12.0; framework Arduino-ESP32 2.0.17 (package 3.20017.241212); GCC Xtensa 8.4.0; esptool 4.9.0; C/C++. Monitor seriale 115200 baud. Toolchain e librerie sono dichiarate in platformio.ini.",
        ),
    ],
)

fill_checklist(
    doc.tables[40],
    [
        (
            "☑",
            "Build riproducibile. ",
            "Dalla root: `pio run -e <environment>`. Per la matrice completa usare `python scripts/build_matrix.py --check`; wrapper Windows/Linux: build_all.bat e build_all.sh. Il 03/08/2026 hanno compilato sensy_2021_V4_white (RAM 65.132 B; flash 1.071.393 B) e sensy_2024_V4_black (RAM 65.044 B; flash 1.035.913 B).",
        ),
        (
            "☑",
            "Configurazioni. ",
            "Otto environment di produzione: 2021 V4 white; 2023 V1 green e V2 black; 2024 V1 green, V2 ENEA, V3 red, V4 green e V4 black. I target 2021/2023 usano ESP32; i 2024 ESP32-S3. Pin, flash/partizioni, PSRAM, relay, LED, UART e feature sono build_flags specifici e fanno parte del contratto di compatibilità.",
        ),
        (
            "◐",
            "Flashing e provisioning. ",
            "Compilare e caricare con `pio run -e <environment> -t upload`, quindi monitorare a 115200 baud. Al primo avvio il firmware inizializza EEPROM, identifica il dispositivo, espone AP di configurazione e riceve SSID/password tramite POST /configWifi. Topic MQTT, versione firmware, low-power e stati relay persistono nel layout EEPROM documentato in FIRMWARE_CONTRACTS.md. Procedura hardware da validare su fixture/porta reale.",
        ),
        (
            "◐",
            "OTA e rollback. ",
            "OTA HTTP controllata ogni 5 cicli e aggiornamento da SD sono attivi. I binari seguono `BT|ST_<board>_V<versione>.bin`, massimo 21 caratteri; scripts/ota_release.py e dashboard locale gestiscono upload, compatibilità, piano e assegnazione. Versione EEPROM cambia solo dopo update riuscito. Rollback automatico non risulta implementato: ritorno a versione precedente richiede assegnazione OTA o flashing di un binario compatibile già approvato.",
        ),
    ],
)

fill_checklist(
    doc.tables[41],
    [
        (
            "☑",
            "Logging. ",
            "Livello di release previsto FW_LOG_LEVEL=1: configurazione mascherata, disponibilità sensori, riepilogo ciclo, esiti rete, warning/errori. FW_LOG_LEVEL=3 aggiunge campioni, JSON e stack watermark; FW_LOG_BILINGUAL=1 abilita traduzioni critiche. Output seriale 115200. Il target sensy_2024_V4_black forza attualmente livello 3: verificare prima della release.",
        ),
        (
            "☑",
            "Watchdog e salute runtime. ",
            "Task FreeRTOS separano monitoring, sniffer e servizi sui core. Reset watchdog protegge scansioni I2C, rete, MQTT e operazioni lunghe; timeout di ciclo forza reboot. Sono presenti controlli heap/stack, mutex I2C, recovery bit-bang delle linee e reinizializzazione hot-plug dei sensori.",
        ),
        (
            "☑",
            "Gestione errori. ",
            "Timeout espliciti su WiFi, MQTT, HTTP, NTP e sensori; retry limitati e fallback. Se rete/MQTT falliscono, dati restano su SPIFFS/SD e vengono eliminati solo dopo pubblicazione riuscita. Ora: server HTTP → NTP → RTC DS1307 → /e.txt. OTA incompleta viene abortita e non aggiorna la versione persistente.",
        ),
        (
            "◐",
            "Debug e known issues. ",
            "Serial RESET, diagnostica sensori all'avvio/ogni 20 cicli, scan I2C e scripts/analyze_soak_log.py supportano troubleshooting. Checklist HARDWARE_VALIDATION.md richiede smoke test e soak ≥200 cicli o 24 h. Test native non eseguito su questa workstation per gcc/g++ assenti; build segnala API close(bool) deprecata in ESPAsyncWebServer/AsyncTCP; qualifica completa resta aperta.",
        ),
    ],
)

artifacts = [
    ("FW-REP-001", "Repository sorgente firmware", "GitHub carminelau/FirmwareSensy", "Software & AI (nome da definire)", "main@7310c42", "☐ Draft  ☑ Reviewed  ☐ Approved"),
    ("FW-CFG-002", "Matrice target, pin e dipendenze", "platformio.ini", "Firmware Lead (da definire)", "working tree 03/08/2026", "☐ Draft  ☑ Reviewed  ☐ Approved"),
    ("FW-ARC-003", "Architettura e contratti compatibilità", "FIRMWARE_OVERVIEW.md; FIRMWARE_CONTRACTS.md", "Firmware Lead (da definire)", "03/08/2026", "☐ Draft  ☑ Reviewed  ☐ Approved"),
    ("FW-VV-004", "Checklist V&V hardware e test contratti", "HARDWARE_VALIDATION.md; test/test_contracts", "Quality + Firmware", "03/08/2026", "☑ Draft  ☐ Reviewed  ☐ Approved"),
    ("FW-REL-005", "Build matrix e processo OTA", "scripts/build_matrix.py; scripts/ota_release.py; tools/ota_dashboard.py", "Release Manager (da definire)", "03/08/2026", "☑ Draft  ☐ Reviewed  ☐ Approved"),
    ("FW-EVD-006", "Build ESP32 + ESP32-S3 rappresentative", ".pio/build; log sessione 03/08/2026", "Firmware Lead (da definire)", "7310c42 + local", "☐ Draft  ☑ Reviewed  ☐ Approved"),
]
table = doc.tables[42]
clone_rows(table, 1 + len(artifacts))
for row, values in zip(table.rows[1:], artifacts):
    for i, value in enumerate(values):
        set_plain_cell(row.cells[i], value, font_size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT)

gaps = [
    ("Qualifica release incompleta", "Ripristinare GCC per test native; eseguire build_matrix --check su 8 target; smoke test e soak ≥200 cicli/24 h con log analizzato.", "Firmware + Quality", "Da definire", "Alta"),
    ("Governance repository non formalizzata", "Nominare owner; proteggere main; imporre review; definire branch/tag/release notes e conservazione immutabile dei binari.", "Technical Lead", "Da definire", "Alta"),
    ("Licenze/SBOM assenti", "Inventariare licenze transitive e fork; generare SBOM e NOTICE; approvare compatibilità d'uso commerciale.", "Firmware + Quality", "Da definire", "Media"),
    ("Rollback automatico non dimostrato", "Definire e provare procedura di ritorno a binario approvato; valutare rollback automatico/validazione boot e criteri pass/fail.", "Firmware Lead", "Da definire", "Alta"),
]
table = doc.tables[43]
clone_rows(table, 1 + len(gaps))
for row, values in zip(table.rows[1:], gaps):
    for i, value in enumerate(values):
        set_plain_cell(row.cells[i], value, font_size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER if i in (3, 4) else WD_ALIGN_PARAGRAPH.LEFT)

# Ask Word to refresh fields (TOC/page references) on open without rewriting them.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.core_properties.title = "PKS Sensy - Sezione 5 Firmware ed edge computing"
doc.core_properties.subject = "Capitolo firmware compilato dalla repository FirmwareSensy"
doc.save(OUTPUT)
print(OUTPUT)
